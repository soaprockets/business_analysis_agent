from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pathlib import Path

output_path = Path("output/project_design_presentation.pptx")
output_path.parent.mkdir(parents=True, exist_ok=True)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color palette
COLOR_TITLE = RGBColor(31, 78, 121)
COLOR_TEXT = RGBColor(51, 51, 51)
COLOR_ACCENT = RGBColor(0, 112, 192)
COLOR_BG = RGBColor(245, 245, 245)


def set_text_frame_style(text_frame, title_font_size=32, body_font_size=20):
    for paragraph in text_frame.paragraphs:
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.space_after = Pt(10)
        for run in paragraph.runs:
            run.font.name = "Microsoft YaHei"
            run.font.color.rgb = COLOR_TEXT
            if paragraph == text_frame.paragraphs[0] and text_frame.paragraphs[0].text:
                run.font.size = Pt(title_font_size)
                run.font.bold = True
                run.font.color.rgb = COLOR_TITLE
            else:
                run.font.size = Pt(body_font_size)


def add_title_slide(prs, title, subtitle):
    slide_layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(slide_layout)

    # Background shape
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_ACCENT
    bg.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(1), Inches(2.2), Inches(11.333), Inches(1.5))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = title
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(44)
    run.font.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)

    sub_box = slide.shapes.add_textbox(Inches(1), Inches(4.0), Inches(11.333), Inches(1))
    tf2 = sub_box.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.alignment = PP_ALIGN.CENTER
    run2 = p2.add_run()
    run2.text = subtitle
    run2.font.name = "Microsoft YaHei"
    run2.font.size = Pt(24)
    run2.font.color.rgb = RGBColor(220, 230, 242)

    return slide


def add_content_slide(prs, title, bullets, note=""):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    # Title bar
    title_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.2))
    title_bar.fill.solid()
    title_bar.fill.fore_color.rgb = COLOR_TITLE
    title_bar.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12.333), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(32)
    run.font.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)

    # Content
    content_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.6), Inches(12), Inches(5.5))
    tf2 = content_box.text_frame
    tf2.word_wrap = True

    for idx, bullet in enumerate(bullets):
        if idx == 0:
            p = tf2.paragraphs[0]
        else:
            p = tf2.add_paragraph()
        p.level = 0
        p.space_after = Pt(14)
        run = p.add_run()
        run.text = f"• {bullet}"
        run.font.name = "Microsoft YaHei"
        run.font.size = Pt(22)
        run.font.color.rgb = COLOR_TEXT

    # Speaker note
    if note:
        notes_slide = slide.notes_slide
        notes_text_frame = notes_slide.notes_text_frame
        notes_text_frame.text = note
        for paragraph in notes_text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.name = "Microsoft YaHei"
                run.font.size = Pt(14)

    return slide


# Slide 1: Title
add_title_slide(
    prs,
    "竞品专家多 Agent 系统",
    "设计思路与关键决策"
)

# Slide 2: Why this project
add_content_slide(
    prs,
    "一、为什么要做这个项目",
    [
        "竞品分析是产品规划和市场战略的核心输入",
        "传统方式依赖人工阅读 PDF / Word / 网页 / 行业报告",
        "信息来源分散、维度不统一、结论难以溯源",
        "网络信息未经验证，容易引入错误",
        "需要自动化、可解释、可验证的竞品分析系统"
    ],
    note="开场说明项目选题的真实背景：企业级软件、SaaS、消费电子等行业都依赖竞品分析，但手工整理效率低、质量不稳定。"
)

# Slide 3: Real pain points
add_content_slide(
    prs,
    "真实痛点",
    [
        "来源分散：PDF、Word、HTML、Markdown、TXT 多种格式",
        "维度不统一：功能、定价、客户案例描述方式各异",
        "事实溯源困难：报告结论缺乏原始出处",
        "网络信息未验证：搜索补充的事实可能与内部资料冲突",
        "分析过程不可复现：不同人员结论差异大"
    ],
    note="对应视频脚本第一步，用五个关键词概括痛点，让观众快速理解问题。"
)

# Slide 4: Solution approach
add_content_slide(
    prs,
    "二、整体解题思路",
    [
        "多 Agent 流水线：每个 Agent 负责单一专业能力",
        "结构化中间产物：Pydantic 数据模型定义 KB / 分析 / 报告",
        "动态维度发现：不预设框架，维度由资料自然演化",
        "事实溯源：每条事实都带 SourceRef 来源引用",
        "人在环路：冲突记录等待人类专家裁决"
    ],
    note="说明系统不是单一大提示词直接出报告，而是拆分成多个可观测、可调试的 Agent。"
)

# Slide 5: Multi-agent architecture
add_content_slide(
    prs,
    "多 Agent 架构",
    [
        "ExpertAgent：读本地文档，抽取产品、维度、事实",
        "SearchAgent：网页搜索，抓取并提取结构化事实",
        "ExpertAgent（二次）：反向验证搜索事实",
        "RAGStore：向量索引，支持检索与验证",
        "AnalysisAgent：对比矩阵、量化评分、SWOT、差距、建议",
        "ReportAgent：综合生成最终报告"
    ],
    note="这是系统核心架构图，建议在录制时切换到 src/ 目录和流程图。"
)

# Slide 6: Dynamic dimensions
add_content_slide(
    prs,
    "关键设计 1：动态维度发现",
    [
        "不预设固定模板",
        "资料谈什么，维度就是什么",
        "避免信息被削足适履",
        "不同产品可以拥有不同维度"
    ],
    note="强调与固定模板的区别：如果资料重点讲 AI 功能和客户案例，这两个就成为维度。"
)

# Slide 7: Fact traceability
add_content_slide(
    prs,
    "关键设计 2：事实溯源",
    [
        "每条事实 = FactItem",
        "包含 content / status / confidence / source_refs",
        "source_refs 记录文档、页码/段落、原文引用",
        "任何结论都可回到原始资料验证"
    ],
    note="展示 models.py 中 FactItem 和 SourceRef 的定义。"
)

# Slide 8: Search & verification
add_content_slide(
    prs,
    "关键设计 3：搜索补充与反向验证",
    [
        "SearchAgent 用 DuckDuckGo / Tavily 搜索网页",
        "LLM 从网页提取结构化事实",
        "ExpertAgent 将事实与原始资料比对",
        "状态：confirmed / external_supplement / conflict",
        "冲突进入 conflicts_and_clarifications 等待人工裁决"
    ],
    note="说明人在环路的作用：不是所有网络信息都直接写进知识库，而是经过验证。"
)

# Slide 9: Analysis vs Report
add_content_slide(
    prs,
    "关键设计 4：分析层与报告层分离",
    [
        "AnalysisAgent 输出结构化结果",
        "对比矩阵 / 评分排名 / SWOT / 差距 / 建议",
        "ReportAgent 基于结构化结果做叙事综合",
        "避免 ReportAgent 重复发明战略结论",
        "JSON 可被其他系统直接消费"
    ],
    note="解释为什么要把分析和报告拆开：减少幻觉、保证一致性、支持机器消费。"
)

# Slide 10: Agent Trace
add_content_slide(
    prs,
    "关键设计 5：Agent 执行过程跟踪",
    [
        "命令行添加 --trace 参数",
        "生成 {kb_id}_trace.md",
        "按 Agent 分章节展示中间产物",
        "可审计、可排查、可证明各 Agent 工作效果"
    ],
    note="建议在录制时运行一次带 --trace 的命令，展示 trace.md 的内容。"
)

# Slide 11: Multi-provider
add_content_slide(
    prs,
    "关键设计 6：多 LLM Provider 支持",
    [
        "Anthropic 官方 API",
        "OpenAI 兼容接口",
        "字节跳动火山方舟 Agent Plan",
        "通过 LLM_PROVIDER 环境变量切换",
        "业务代码几乎无需改动"
    ],
    note="展示 .env 中 volcano-agent-plan 的配置方式。"
)

# Slide 12: Outputs
add_content_slide(
    prs,
    "运行效果与输出",
    [
        "kb_xxxxxx.json / .md：结构化知识库",
        "kb_xxxxxx_report.json / .md：最终竞品分析报告",
        "kb_xxxxxx_trace.md：Agent 执行跟踪",
        "示例：CloudFlow CRM vs SalesMax Pro",
        "自动生成评分、排名、SWOT、战略建议"
    ],
    note="切换到 output/ 目录，展示生成的所有文件。"
)

# Slide 13: Summary
add_content_slide(
    prs,
    "总结",
    [
        "多 Agent 流水线替代单一大模型提示",
        "动态维度替代固定模板",
        "事实溯源 + 反向验证保证可信度",
        "结构化 JSON + Markdown 双输出",
        "Agent Trace 让分析过程透明可审计"
    ],
    note="收尾，强调项目的核心创新点和实际价值。"
)

prs.save(output_path)
print(f"已生成 PPT：{output_path.resolve()}")
