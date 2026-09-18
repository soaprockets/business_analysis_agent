"""
Document parsers: extract plain text from various readable file formats.
"""

import os
import re
import uuid
from pathlib import Path
from typing import Optional

from bs4 import BeautifulSoup


class ParsedDocument:
    """Result of parsing a single document."""

    def __init__(
        self,
        doc_id: str,
        file_name: str,
        file_type: str,
        text: str,
        pages: Optional[list[str]] = None,
    ):
        self.doc_id = doc_id
        self.file_name = file_name
        self.file_type = file_type
        self.text = text
        self.pages = pages or []

    @property
    def chunks(self) -> list[str]:
        """Return text split into manageable overlapping chunks."""
        return chunk_text(self.text)


def chunk_text(
    text: str, chunk_size: int = 4000, overlap: int = 500
) -> list[str]:
    """Split text into overlapping chunks by paragraph boundaries."""
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    chunks: list[str] = []
    current_chunk: list[str] = []
    current_length = 0

    for paragraph in paragraphs:
        para_len = len(paragraph)
        if current_length + para_len > chunk_size and current_chunk:
            chunks.append("\n\n".join(current_chunk))
            # Keep overlap paragraphs
            overlap_text = []
            overlap_len = 0
            for p in reversed(current_chunk):
                if overlap_len + len(p) > overlap:
                    break
                overlap_text.insert(0, p)
                overlap_len += len(p)
            current_chunk = overlap_text
            current_length = overlap_len

        current_chunk.append(paragraph)
        current_length += para_len

    if current_chunk:
        chunks.append("\n\n".join(current_chunk))

    return chunks if chunks else [text]


def parse_document(file_path: str | Path) -> ParsedDocument:
    """Route to the correct parser based on file extension."""
    file_path = Path(file_path)
    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    file_name = file_path.name
    file_type = file_path.suffix.lower()
    doc_id = f"doc_{uuid.uuid4().hex[:8]}"

    if file_type == ".pdf":
        return _parse_pdf(file_path, doc_id, file_name)
    if file_type in {".docx", ".doc"}:
        return _parse_docx(file_path, doc_id, file_name)
    if file_type in {".html", ".htm"}:
        return _parse_html(file_path, doc_id, file_name)
    if file_type in {".md", ".markdown", ".txt", ".json", ".csv"}:
        return _parse_text(file_path, doc_id, file_name)

    # Fallback: try to read as plain text
    try:
        return _parse_text(file_path, doc_id, file_name)
    except Exception as exc:
        raise ValueError(f"Unsupported file type: {file_type}") from exc


def _parse_pdf(file_path: Path, doc_id: str, file_name: str) -> ParsedDocument:
    from pypdf import PdfReader

    reader = PdfReader(str(file_path))
    pages: list[str] = []
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            pages.append(f"--- Page {i} ---\n{text.strip()}")

    full_text = "\n\n".join(pages)
    return ParsedDocument(doc_id, file_name, ".pdf", full_text, pages)


def _parse_docx(file_path: Path, doc_id: str, file_name: str) -> ParsedDocument:
    from docx import Document

    doc = Document(str(file_path))
    paragraphs: list[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            paragraphs.append(para.text.strip())

    full_text = "\n\n".join(paragraphs)
    return ParsedDocument(doc_id, file_name, ".docx", full_text)


def _parse_html(file_path: Path, doc_id: str, file_name: str) -> ParsedDocument:
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        soup = BeautifulSoup(f.read(), "html.parser")
    for tag in soup(["script", "style"]):
        tag.decompose()
    text = soup.get_text(separator="\n", strip=True)
    return ParsedDocument(doc_id, file_name, ".html", text)


def _parse_text(file_path: Path, doc_id: str, file_name: str) -> ParsedDocument:
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        text = f.read()
    return ParsedDocument(doc_id, file_name, file_path.suffix.lower(), text)


def parse_documents(file_paths: list[str | Path]) -> list[ParsedDocument]:
    """Parse multiple documents."""
    return [parse_document(fp) for fp in file_paths]
